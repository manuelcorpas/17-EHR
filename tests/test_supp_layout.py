import subprocess, tempfile
from importlib.util import spec_from_file_location, module_from_spec
from pathlib import Path
import docx
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
spec = spec_from_file_location('st', ROOT / 'PYTHON' / '08-03-supplementary-tables.py'); st = module_from_spec(spec); spec.loader.exec_module(st)


def _render():
    d = Path(tempfile.mkdtemp()); md = d / 's.md'; out = d / 's.docx'
    md.write_text(st.build()); subprocess.run(['pandoc', str(md), '-o', str(out)], check=True)
    st.format_docx(out)
    return docx.Document(str(out))


def test_supplement_tables_are_laid_out_for_reading_in_portrait():
    doc = _render()
    for s in doc.sections:
        assert s.orientation == WD_ORIENT.PORTRAIT and s.page_width < s.page_height
        assert s.left_margin <= Pt(54) and s.right_margin <= Pt(54)
    big = [t for t in doc.tables if len(t.rows) > 100]
    assert len(big) == 2
    for t in big:
        r0, r1 = t.rows[0], t.rows[1]
        assert r0._tr.trPr is not None and r0._tr.trPr.find(docx.oxml.ns.qn('w:tblHeader')) is not None    # header repeats
        assert r1._tr.trPr.find(docx.oxml.ns.qn('w:tblHeader')) is not None
        assert r1.cells[1].text == 'Before' and r1.cells[2].text == 'After'                           # grouped header
        widths = [c.width for c in t.rows[2].cells]
        assert widths[0] == max(widths) and len(set(widths)) > 1                                              # condition column widest
        assert sum(widths) <= doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
        sizes = {r.font.size for c in t.rows[5].cells for p in c.paragraphs for r in p.runs if r.text.strip()}
        assert sizes == {Pt(8)}
        num = t.rows[5].cells[3].paragraphs[0]
        assert num.alignment == WD_ALIGN_PARAGRAPH.RIGHT


def test_layout_does_not_change_table_data():
    d = Path(tempfile.mkdtemp()); md = d / 's.md'; out = d / 's.docx'
    md.write_text(st.build()); subprocess.run(['pandoc', str(md), '-o', str(out)], check=True)
    rows = lambda t, skip: [[c.text for c in r.cells] for r in t.rows[skip:]]
    before = docx.Document(str(out)).tables
    st.format_docx(out); after = docx.Document(str(out)).tables
    assert len(before) == len(after)
    for b, a in zip(before, after):
        grouped = len(b.rows) > 100
        assert rows(a, 2 if grouped else 1) == rows(b, 1)            # data rows untouched


def test_counts_and_ranks_print_as_integers():
    md = st.build()
    assert '| 89.0 |' not in md and '| 5,229.0 |' not in md and '| 1041 |' not in md
    assert '| 5,229 |' in md and '| 1,041 |' in md


def test_header_has_no_vertical_merges():
    doc = _render()
    for t in [t for t in doc.tables if len(t.rows) > 100]:
        for r in t.rows[:2]:
            assert r._tr.xpath('.//w:vMerge') == [] and len(r._tr.xpath('./w:trPr/w:tblHeader')) == 1
        assert t.rows[1].cells[0].text == 'Condition'
